"""Text aus Dateien auslesen, z. B. aus Mail-Anhängen (inbox.py).

  PDF                      pypdf (nur PDFs mit echtem Text; eingescannte enthalten nur Bilder)
  Word / PowerPoint        .docx, .pptx - ZIP mit XML, ausgelesen mit der Standardbibliothek
  Excel                    .xlsx, .xlsm - openpyxl
  HTML, E-Mail (.eml)      HTML wird zu Text, eine angehängte Mail liefert ihren Text
  Textdateien              .txt, .csv, .json, .ics … und alles mit text/* als Typ

Andere Dateien (Bilder, Archive, Programme …) liefern None: Sie lassen sich im Chat
öffnen oder speichern, aber nicht als Text lesen.

Die Dateien kommen von Fremden, deshalb Grenzen gegen präparierte Dateien: höchstens
MAX_PDF_PAGES Seiten und MAX_UNPACKED_BYTES entpackt (gegen "ZIP-Bomben").
"""
from __future__ import annotations

import email
import email.policy
import io
import re
import zipfile
from email.message import Message
from html.parser import HTMLParser
from pathlib import PurePath
from xml.etree import ElementTree

MAX_PDF_PAGES = 100
MAX_SHEET_ROWS = 2_000
MAX_UNPACKED_BYTES = 50 * 1024 * 1024

TEXT_SUFFIXES = frozenset({".txt", ".csv", ".tsv", ".md", ".json", ".xml", ".log", ".ics", ".vcf",
                           ".ini", ".cfg", ".yaml", ".yml", ".py", ".js", ".java", ".c", ".sql"})

_WORD = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_DRAWING = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


class FileTextError(ValueError):
    """Datei ist ein bekannter Typ, lässt sich aber nicht lesen (defekt, verschlüsselt …)."""


def extract_text(data: bytes, filename: str, mime_type: str = "") -> str | None:
    """Text der Datei; None, wenn der Dateityp keinen lesbaren Text hat."""
    readers = {"pdf": _pdf_text, "docx": _docx_text, "pptx": _pptx_text, "xlsx": _xlsx_text,
               "html": lambda d: html_to_text(decode_text(d)),
               "eml": lambda d: mail_body_text(email.message_from_bytes(d, policy=email.policy.default)),
               "text": decode_text}
    kind = _kind(filename, mime_type)
    return readers[kind](data) if kind else None


def can_extract(filename: str, mime_type: str = "") -> bool:
    return _kind(filename, mime_type) is not None


def _kind(filename: str, mime_type: str) -> str | None:
    suffix = PurePath(filename).suffix.lower()
    mime_type = mime_type.lower()
    if suffix == ".pdf" or mime_type == "application/pdf":
        return "pdf"
    if suffix in (".docx", ".pptx"):
        return suffix[1:]
    if suffix in (".xlsx", ".xlsm"):
        return "xlsx"
    if suffix in (".html", ".htm") or mime_type == "text/html":
        return "html"
    if suffix == ".eml" or mime_type == "message/rfc822":
        return "eml"
    if suffix in TEXT_SUFFIXES or mime_type.startswith("text/"):
        return "text"
    return None


def decode_text(data: bytes) -> str:
    """UTF-8, sonst Windows-1252 (ältere Windows-Programme, z. B. Excel-CSV)."""
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def mail_body_text(message: Message) -> str:
    """Text einer Mail: text/plain, sonst text/html als Text; Leerzeilen zusammengefasst."""
    part = message.get_body(preferencelist=("plain", "html"))
    if part is None:
        return ""
    try:
        text = part.get_content()
    except (LookupError, UnicodeError):     # unbekannter oder falsch angegebener Zeichensatz
        text = decode_text(part.get_payload(decode=True) or b"")
    if part.get_content_subtype() == "html":
        text = html_to_text(text)
    return tidy(text)


def tidy(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


# --- PDF ---------------------------------------------------------------------------

def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader          # erst hier: Ohne pypdf funktionieren die anderen Typen trotzdem
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise FileTextError("PDF ist mit einem Passwort geschützt")
        pages = []
        for number, page in enumerate(reader.pages[:MAX_PDF_PAGES], start=1):
            text = tidy(page.extract_text() or "")
            if text:
                pages.append(f"[Seite {number}]\n{text}")
        total = len(reader.pages)
    except FileTextError:
        raise
    except Exception as error:   # noqa: BLE001  # pypdf wirft bei defekten PDFs viele verschiedene Fehler
        raise FileTextError(f"PDF konnte nicht gelesen werden: {error}") from None
    if not pages:
        raise FileTextError("PDF enthält keinen Text (vermutlich eingescannt - nur als Bild vorhanden)")
    if total > MAX_PDF_PAGES:
        pages.append(f"[... nur die ersten {MAX_PDF_PAGES} von {total} Seiten gelesen]")
    return "\n\n".join(pages)


# --- Office (ZIP mit XML) ----------------------------------------------------------

def _open_zip(data: bytes, kind: str) -> zipfile.ZipFile:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise FileTextError(f"{kind}-Datei ist beschädigt") from None
    if sum(info.file_size for info in archive.infolist()) > MAX_UNPACKED_BYTES:
        raise FileTextError(f"{kind}-Datei ist entpackt zu groß")
    return archive


def _xml(archive: zipfile.ZipFile, name: str, kind: str) -> ElementTree.Element:
    try:
        return ElementTree.fromstring(archive.read(name))
    except (KeyError, ElementTree.ParseError):
        raise FileTextError(f"{kind}-Datei ist beschädigt ({name})") from None


def _paragraphs(root: ElementTree.Element, namespace: str) -> list[str]:
    paragraphs = []
    for paragraph in root.iter(f"{namespace}p"):
        parts = []
        for node in paragraph.iter():
            if node.tag == f"{namespace}t":
                parts.append(node.text or "")
            elif node.tag == f"{namespace}tab":
                parts.append("\t")
            elif node.tag in (f"{namespace}br", f"{namespace}cr"):
                parts.append("\n")
        paragraphs.append("".join(parts))
    return paragraphs


def _docx_text(data: bytes) -> str:
    with _open_zip(data, "Word") as archive:
        return tidy("\n".join(_paragraphs(_xml(archive, "word/document.xml", "Word"), _WORD)))


def _pptx_text(data: bytes) -> str:
    with _open_zip(data, "PowerPoint") as archive:
        slides = sorted((name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
                        key=lambda name: int(re.search(r"\d+", name).group()))
        parts = []
        for number, name in enumerate(slides, start=1):
            text = tidy("\n".join(_paragraphs(_xml(archive, name, "PowerPoint"), _DRAWING)))
            parts.append(f"[Folie {number}]\n{text}")
    return "\n\n".join(parts)


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook
    _open_zip(data, "Excel").close()     # Größe prüfen, bevor openpyxl entpackt
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as error:   # noqa: BLE001  # openpyxl meldet defekte Dateien uneinheitlich
        raise FileTextError(f"Excel-Datei konnte nicht gelesen werden: {error}") from None
    sheets = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value) for value in row]
                if any(cells):
                    rows.append(" | ".join(cells).rstrip(" |"))
                if len(rows) >= MAX_SHEET_ROWS:
                    rows.append(f"[... nach {MAX_SHEET_ROWS} Zeilen abgebrochen]")
                    break
            sheets.append(f"[Tabellenblatt {sheet.title}]\n" + ("\n".join(rows) or "(leer)"))
    finally:
        workbook.close()
    return "\n\n".join(sheets)


# --- HTML --------------------------------------------------------------------------

class _TextExtractor(HTMLParser):
    BLOCKS = {"p", "div", "br", "tr", "li", "table", "blockquote", "h1", "h2", "h3", "h4", "h5", "h6"}
    HIDDEN = {"script", "style", "head"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.HIDDEN:
            self.hidden += 1
        elif tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.HIDDEN:
            self.hidden = max(0, self.hidden - 1)
        elif tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def html_to_text(html: str) -> str:
    parser = _TextExtractor()
    parser.feed(html)
    parser.close()
    return tidy("\n".join(" ".join(line.split()) for line in "".join(parser.parts).split("\n")))
